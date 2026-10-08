import glob
import io
import os
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from wifi_diag import pdf, report
from wifi_diag.backends import DemoBackend
from wifi_diag.model import Snapshot
from wifi_diag.rules import analyze, channel_utilization
from wifi_diag.web.server import make_server
from wifi_diag.web.state import AppState
from wifi_diag.xlsx import build_xlsx

from ie_builder import *  # noqa: F401,F403
from test_web import Client

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

CHROME = next(iter(glob.glob("/opt/pw-browsers/chromium-*/chrome-linux/chrome")), None) or shutil.which("chromium")


@unittest.skipIf(openpyxl is None, "нужен openpyxl для проверки")
class XlsxTests(unittest.TestCase):
    def test_writer_roundtrip_and_hostile_text(self):
        evil = "SSID\x00\x07<b>&\"'\ud800 Привет"
        data = build_xlsx([("Лист: [1]", [["Имя", "Число", "Дробь", "Флаг", "Пусто"], [evil, 42, -61.5, True, None]]),
                           ("Лист: [1]", [["x"]])])
        wb = openpyxl.load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames, ["Лист_ _1_", "Лист_ _1_ (2)"])
        ws = wb.worksheets[0]
        self.assertEqual([c.value for c in ws[1]], ["Имя", "Число", "Дробь", "Флаг", "Пусто"])
        self.assertEqual([c.value for c in ws[2]][:4], ["SSID<b>&\"' Привет", 42, -61.5, "да"])
        self.assertTrue(ws["A1"].font.b)
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws.auto_filter.ref, "A1:E2")

    def test_report_workbook(self):
        bss = [make_bss("02:00:00:00:20:01", "CORP", CH36, -55, rsn_ie([1]) + qbss_ie(10, 180)),
               make_bss("02:00:00:00:20:02", "Neighbor", CH6, -70, rsn_ie([2]))]
        snap = Snapshot("2026-10-08T12:00:00+03:00", "Adapter", bss)
        f = analyze(bss, focus_ssids=["CORP"])
        journal = {"events": [{"t": 1e9, "kind": "roam", "severity": "info", "text": "Роуминг", "ssid": "CORP",
                               "bssid_from": "a", "bssid_to": "b", "rssi_from": -70, "rssi_to": -50, "duration_s": 12.0,
                               "reason_code": None, "reason": ""}], "stats": {"roams": 1}}
        survey = [{"id": 1, "x": 0.5, "y": 0.25, "t": 1e9, "status": "done", "conn": {"ssid": "CORP", "bssid": "a", "rssi": -60},
                   "results": {"02:00:00:00:20:01": ["CORP", "5", 36, -58], "x": ["CORP", "2.4", 1, -80]}}]
        data = report.build_xlsx_report(snap, f, ["CORP"], "только мои", util=channel_utilization(bss),
                                        history={"02:00:00:00:20:01": [(1e9, -55, 70.6)]}, journal=journal, survey=survey)
        wb = openpyxl.load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames, ["Сводка", "BSS", "Находки", "Загрузка каналов", "История", "Журнал подключения", "Обход"])
        self.assertEqual(wb["BSS"].max_row, 3)
        self.assertEqual(wb["Загрузка каналов"]["C2"].value, 70.6)
        obh = [c.value for c in wb["Обход"][2]]
        self.assertEqual(obh[0:3], [1, 50.0, 25.0])
        self.assertEqual(obh[-2:], [-58, 1])           # лучший RSSI CORP и сколько BSS CORP от -75

    @unittest.skipIf(shutil.which("soffice") is None, "нет LibreOffice")
    def test_libreoffice_opens_it(self):
        data = build_xlsx([("BSS", [["SSID", "RSSI"], ["Сеть", -60]])])
        with tempfile.TemporaryDirectory() as d:
            src = os.path.join(d, "r.xlsx")
            with open(src, "wb") as fh:
                fh.write(data)
            subprocess.run(["soffice", "--headless", "--convert-to", "csv", "--outdir", d, src], timeout=120,
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            with open(os.path.join(d, "r.csv"), encoding="utf-8") as fh:
                self.assertEqual(fh.read().split(), ["SSID,RSSI", "Сеть,-60"])


class ExportApiTests(unittest.TestCase):
    def setUp(self):
        self.state = AppState(DemoBackend(seed=4), None, ["CORP"])
        self.server = make_server(self.state)
        threading.Thread(target=self.server.serve_forever, daemon=True).start()
        self.state.start_monitors()
        self.c = Client(self.server, self.server.token)
        self.c.j("POST", "/api/scan", {"interface": 0, "wait": 1})
        for _ in range(100):
            time.sleep(0.05)
            if self.state.scan_count and not self.state.scanning:
                break

    def tearDown(self):
        self.state.stop_monitors()
        self.server.shutdown()
        self.server.server_close()

    @unittest.skipIf(openpyxl is None, "нужен openpyxl")
    def test_xlsx_export_respects_scope(self):
        st, hd, raw = self.c.req("POST", "/api/export/xlsx", {"scope": "focus"})
        self.assertEqual(st, 200)
        self.assertIn("wifi-report-my.xlsx", hd["Content-Disposition"])
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        self.assertEqual({r[0] for r in wb["BSS"].iter_rows(min_row=2, values_only=True)}, {"CORP"})
        self.assertIn("Журнал подключения", wb.sheetnames)

    @unittest.skipIf(not CHROME or shutil.which("pdftotext") is None, "нет Chromium или pdftotext")
    def test_pdf_export(self):
        with mock.patch.dict(os.environ, {"WIFI_DIAG_PDF_BROWSER": CHROME}):
            st, hd, raw = self.c.req("POST", "/api/export/pdf", {"scope": "all"})
        self.assertEqual(st, 200, raw[:200])
        self.assertTrue(raw.startswith(b"%PDF"))
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "r.pdf")
            with open(p, "wb") as fh:
                fh.write(raw)
            text = subprocess.run(["pdftotext", "-layout", p, "-"], capture_output=True, timeout=60).stdout.decode("utf-8")
        for needle in ("Wi-Fi: отчёт диагностики", "Загрузка каналов", "Подключение ноутбука", "Все BSS", "CORP"):
            self.assertIn(needle, text)

    def test_pdf_without_browser_gives_clear_error(self):
        with mock.patch.object(pdf, "find_browser", return_value=None):
            st, body = self.c.j("POST", "/api/export/pdf", {"scope": "all"})
        self.assertEqual(st, 500)
        self.assertIn("Edge", body["error"])

    def test_html_report_has_new_sections_and_escapes(self):
        st, _, raw = self.c.req("POST", "/api/export/html", {"scope": "all"})
        h = raw.decode("utf-8")
        self.assertIn("Загрузка каналов (QBSS Load)", h)
        self.assertIn("Подключение ноутбука", h)
        self.assertIn("@page", h)
        self.assertIn("default-src 'none'", h)


if __name__ == "__main__":
    unittest.main()
